"""
Testes de fluxo para versionamento de planos (XLSX local, sem Google Sheets).

Estratégia: funções Python puras com arquivos reais em tmp_path —
nenhum mock de lógica, só redirecionamos o diretório de versões.
"""
from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook

from core.versioning import (
    delete_plano_version_xlsx,
    list_plano_versions_xlsx,
    restore_plano_version_xlsx,
    save_plano_version_xlsx,
)


# ── fixture ───────────────────────────────────────────────────────────────────

def _make_plano(path: Path, seed: str = "SKU001") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plano_Enderecamento_Final"
    ws.append(["location_id", "product_code", "product_name"])
    ws.append(["R1E1-A1", seed, f"Produto {seed}"])
    wb.save(path)
    return path


@pytest.fixture()
def plano(tmp_path):
    return _make_plano(tmp_path / "plano.xlsx")


@pytest.fixture()
def ver_dir(tmp_path):
    return tmp_path / "versions"


# ── salvar versão ─────────────────────────────────────────────────────────────

class TestSalvarVersao:
    def test_retorna_success_true(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = save_plano_version_xlsx(plano, "v1")
        assert r["success"] is True

    def test_retorna_version_id(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = save_plano_version_xlsx(plano, "v1")
        assert r.get("version_id")

    def test_cria_arquivo_fisico(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = save_plano_version_xlsx(plano, "v1")
        assert (ver_dir / r["version_id"]).exists()

    def test_duas_versoes_geram_arquivos_distintos(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r1 = save_plano_version_xlsx(plano, "primeira")
            r2 = save_plano_version_xlsx(plano, "segunda")
        assert r1["version_id"] != r2["version_id"]

    def test_arquivo_inexistente_retorna_error(self, tmp_path, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = save_plano_version_xlsx(tmp_path / "ghost.xlsx", "v1")
        assert r["success"] is False


# ── listar versões ────────────────────────────────────────────────────────────

class TestListarVersoes:
    def test_sem_versoes_retorna_lista_vazia(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = list_plano_versions_xlsx(plano)
        assert r["success"] is True
        assert r["versions"] == []

    def test_apos_salvar_aparece_na_lista(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            save_plano_version_xlsx(plano, "v1")
            r = list_plano_versions_xlsx(plano)
        assert len(r["versions"]) == 1

    def test_duas_versoes_aparecem_na_lista(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            save_plano_version_xlsx(plano, "v1")
            save_plano_version_xlsx(plano, "v2")
            r = list_plano_versions_xlsx(plano)
        assert len(r["versions"]) == 2

    def test_cada_versao_tem_version_id_e_label(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            save_plano_version_xlsx(plano, "v1")
            r = list_plano_versions_xlsx(plano)
        v = r["versions"][0]
        assert "version_id" in v and "label" in v

    def test_ordenadas_mais_recente_primeiro(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            save_plano_version_xlsx(plano, "antiga")
            save_plano_version_xlsx(plano, "nova")
            r = list_plano_versions_xlsx(plano)
        ts = [v["timestamp"] for v in r["versions"]]
        assert ts == sorted(ts, reverse=True)


# ── restaurar versão ──────────────────────────────────────────────────────────

class TestRestaurarVersao:
    def test_restaurar_retorna_success(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r_save = save_plano_version_xlsx(plano, "snap")
            r = restore_plano_version_xlsx(plano, r_save["version_id"])
        assert r["success"] is True

    def test_dados_originais_sao_recuperados(self, tmp_path, ver_dir):
        plano = _make_plano(tmp_path / "plano.xlsx", seed="ORIGINAL")
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r_save = save_plano_version_xlsx(plano, "snap")
            _make_plano(plano, seed="MODIFICADO")
            restore_plano_version_xlsx(plano, r_save["version_id"])

        wb = load_workbook(plano)
        ws = wb["Plano_Enderecamento_Final"]
        row2_code = ws.cell(row=2, column=2).value
        assert row2_code == "ORIGINAL"

    def test_versao_inexistente_retorna_error(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = restore_plano_version_xlsx(plano, "nao_existe.xlsx")
        assert r["success"] is False

    def test_arquivo_base_inexistente_retorna_error(self, tmp_path, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = restore_plano_version_xlsx(tmp_path / "ghost.xlsx", "v1.xlsx")
        assert r["success"] is False


# ── deletar versão ────────────────────────────────────────────────────────────

class TestDeletarVersao:
    def test_deletar_retorna_success(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r_save = save_plano_version_xlsx(plano, "para_deletar")
            r = delete_plano_version_xlsx(r_save["version_id"])
        assert r["success"] is True

    def test_arquivo_e_removido_do_disco(self, plano, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r_save = save_plano_version_xlsx(plano, "para_deletar")
            version_id = r_save["version_id"]
            delete_plano_version_xlsx(version_id)
        assert not (ver_dir / version_id).exists()

    def test_versao_inexistente_retorna_error(self, ver_dir):
        with patch("core.versioning.VERSION_DIR", ver_dir):
            r = delete_plano_version_xlsx("nao_existe.xlsx")
        assert r["success"] is False


# ── fluxo completo ────────────────────────────────────────────────────────────

class TestFluxoCompletoVersoes:
    def test_salvar_listar_restaurar_deletar(self, tmp_path, ver_dir):
        """
        Simula o ciclo real do operador:
        salvar snapshot → listar → restaurar dados → deletar snapshot.
        """
        plano = _make_plano(tmp_path / "plano.xlsx", seed="PRODUCAO")
        with patch("core.versioning.VERSION_DIR", ver_dir):
            # 1. Salvar
            r_save = save_plano_version_xlsx(plano, "producao")
            assert r_save["success"] is True
            vid = r_save["version_id"]

            # 2. Listar — deve aparecer
            r_list = list_plano_versions_xlsx(plano)
            assert any(v["version_id"] == vid for v in r_list["versions"])

            # 3. Simular edição e restaurar
            _make_plano(plano, seed="EDITADO")
            r_restore = restore_plano_version_xlsx(plano, vid)
            assert r_restore["success"] is True

            # Dados de volta ao estado original
            wb = load_workbook(plano)
            assert wb["Plano_Enderecamento_Final"].cell(row=2, column=2).value == "PRODUCAO"

            # 4. Deletar
            r_del = delete_plano_version_xlsx(vid)
            assert r_del["success"] is True

            # 5. Listar — não deve mais aparecer
            r_list2 = list_plano_versions_xlsx(plano)
            assert not any(v["version_id"] == vid for v in r_list2["versions"])
