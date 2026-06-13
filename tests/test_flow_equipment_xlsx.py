"""
Testes de fluxo para troca de equipamentos e remoção em massa (XLSX local).

Estratégia: funções Python puras com arquivos reais em tmp_path.
Verifica que os dados no disco mudam como esperado após cada operação.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.moves import execute_equipment_swap
from core.bulk_remove import (
    preview_remove_all_products_by_filter_xlsx,
    remove_all_products_by_filter_xlsx,
)


# ── fixtures ──────────────────────────────────────────────────────────────────

def _make_swap_xlsx(path: Path) -> Path:
    """
    Dois equipamentos de prateleira no mesmo arquivo:
    - R1E1: slot A1 com SKU001
    - R1E2: slot A1 com SKU002
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Plano_Enderecamento_Final"
    ws.append(["rua_num", "equipamento_num", "tipo_equipamento", "location_id", "product_code", "product_name"])
    ws.append([1, 1, "prateleira", "R1E1-A1", "SKU001", "Produto A"])
    ws.append([1, 2, "prateleira", "R1E2-A1", "SKU002", "Produto B"])
    wb.save(path)
    return path


def _make_bulk_xlsx(path: Path) -> Path:
    """
    Base_Produtos com um produto quimico e um alimento.
    Plano com ambos alocados.
    """
    wb = Workbook()

    ws_base = wb.active
    ws_base.title = "Base_Produtos"
    ws_base.append(["product_code", "produto", "grupo", "categoria_armazenagem"])
    ws_base.append(["SKU-QUIM", "Produto Quimico", "quimico", "seco"])
    ws_base.append(["SKU-ALIM", "Produto Alimento", "alimento", "seco"])

    ws_plano = wb.create_sheet("Plano_Enderecamento_Final")
    ws_plano.append(["location_id", "product_code", "product_name"])
    ws_plano.append(["R1E1-A1", "SKU-QUIM", "Produto Quimico"])
    ws_plano.append(["R1E2-A1", "SKU-ALIM", "Produto Alimento"])

    wb.save(path)
    return path


def _read_product_codes(path: Path) -> list[str | None]:
    wb = load_workbook(path)
    ws = wb["Plano_Enderecamento_Final"]
    headers = [c.value for c in ws[1]]
    col = headers.index("product_code") + 1
    return [ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)]


# ── troca de equipamentos ─────────────────────────────────────────────────────

class TestTrocaEquipamentos:
    def test_troca_retorna_success(self, tmp_path):
        xlsx = _make_swap_xlsx(tmp_path / "plano.xlsx")
        r = execute_equipment_swap(xlsx, "R1E1", "R1E2")
        assert r["success"] is True

    def test_produtos_ficam_invertidos_apos_troca(self, tmp_path):
        xlsx = _make_swap_xlsx(tmp_path / "plano.xlsx")
        execute_equipment_swap(xlsx, "R1E1", "R1E2")

        codes = _read_product_codes(xlsx)
        # R1E1 (linha 2) deve ter SKU002; R1E2 (linha 3) deve ter SKU001
        assert codes[0] == "SKU002"
        assert codes[1] == "SKU001"

    def test_troca_dupla_volta_ao_estado_original(self, tmp_path):
        xlsx = _make_swap_xlsx(tmp_path / "plano.xlsx")
        execute_equipment_swap(xlsx, "R1E1", "R1E2")
        execute_equipment_swap(xlsx, "R1E1", "R1E2")

        codes = _read_product_codes(xlsx)
        assert codes[0] == "SKU001"
        assert codes[1] == "SKU002"

    def test_tipos_incompativeis_retorna_error(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Plano_Enderecamento_Final"
        ws.append(["rua_num", "equipamento_num", "tipo_equipamento", "location_id", "product_code"])
        ws.append([1, 1, "prateleira", "R1E1-A1", "SKU001"])
        ws.append([1, 2, "geladeira",  "R1E2-A1", "SKU002"])
        path = tmp_path / "plano.xlsx"
        wb.save(path)

        r = execute_equipment_swap(path, "R1E1", "R1E2")
        assert r["success"] is False
        assert "incompatív" in r.get("error", "").lower() or "incompativel" in r.get("error", "").lower()

    def test_id_invalido_retorna_error(self, tmp_path):
        xlsx = _make_swap_xlsx(tmp_path / "plano.xlsx")
        r = execute_equipment_swap(xlsx, "INVALIDO", "R1E2")
        assert r["success"] is False

    def test_equipamento_inexistente_retorna_error(self, tmp_path):
        xlsx = _make_swap_xlsx(tmp_path / "plano.xlsx")
        r = execute_equipment_swap(xlsx, "R9E9", "R8E8")
        assert r["success"] is False

    def test_aba_ausente_retorna_error(self, tmp_path):
        wb = Workbook()
        wb.active.title = "Outra_Aba"
        path = tmp_path / "plano.xlsx"
        wb.save(path)
        r = execute_equipment_swap(path, "R1E1", "R1E2")
        assert r["success"] is False

    def test_colunas_essenciais_ausentes_retorna_error(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Plano_Enderecamento_Final"
        ws.append(["apenas_uma_coluna"])
        ws.append(["valor"])
        path = tmp_path / "plano.xlsx"
        wb.save(path)
        r = execute_equipment_swap(path, "R1E1", "R1E2")
        assert r["success"] is False


# ── remoção em massa ──────────────────────────────────────────────────────────

class TestPreviewRemocao:
    def test_conta_skus_afetados(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = preview_remove_all_products_by_filter_xlsx(xlsx, "quimicos")
        assert r["success"] is True
        assert r["sku_count"] == 1
        assert r["plano_count"] == 1

    def test_nao_modifica_arquivo(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        codes_antes = _read_product_codes(xlsx)
        preview_remove_all_products_by_filter_xlsx(xlsx, "quimicos")
        assert _read_product_codes(xlsx) == codes_antes

    def test_filtro_sem_matches_retorna_zero(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = preview_remove_all_products_by_filter_xlsx(xlsx, "freezer")
        assert r["success"] is True
        assert r["sku_count"] == 0

    def test_filtro_invalido_retorna_error(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = preview_remove_all_products_by_filter_xlsx(xlsx, "filtro_inexistente")
        assert r["success"] is False


class TestRemocaoEfetiva:
    def test_remove_produto_quimico_do_plano(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = remove_all_products_by_filter_xlsx(xlsx, "quimicos")
        assert r["success"] is True
        assert r["plano_updated"] == 1

    def test_produto_de_outro_grupo_permanece(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        remove_all_products_by_filter_xlsx(xlsx, "quimicos")
        codes = _read_product_codes(xlsx)
        assert "SKU-ALIM" in codes

    def test_produto_removido_nao_aparece_mais(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        remove_all_products_by_filter_xlsx(xlsx, "quimicos")
        codes = [c for c in _read_product_codes(xlsx) if c]
        assert "SKU-QUIM" not in codes

    def test_filtro_sem_matches_retorna_zero_atualizados(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = remove_all_products_by_filter_xlsx(xlsx, "freezer")
        assert r["success"] is True
        assert r["plano_updated"] == 0

    def test_filtro_invalido_retorna_error(self, tmp_path):
        xlsx = _make_bulk_xlsx(tmp_path / "plano.xlsx")
        r = remove_all_products_by_filter_xlsx(xlsx, "filtro_invalido")
        assert r["success"] is False
