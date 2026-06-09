from copy import copy
from collections import defaultdict
from pathlib import Path

import openpyxl


INPUT_PATH = Path("/Users/andrelobo/Downloads/RevisaoFinal_corrigida_v2.xlsx")
SKELETON_PATH = Path("/Users/andrelobo/Downloads/enderecos.xlsx")
JOIN_PATH = Path("/Users/andrelobo/Downloads/Plano_Enderecamento_Final_JOIN_ATUALIZADO.xlsx")
OUTPUT_PATH = Path("/Users/andrelobo/Downloads/RevisaoFinal_corrigida_v3.xlsx")

SKELETON_SHEET = "Página1"
JOIN_SHEET = "Plano_Enderecamento_Final_JOIN"
PLAN_SHEET = "Plano_Enderecamento_Final"

PRODUCT_COLUMNS = [
    "product_code",
    "product_name",
    "quantidade",
    "curva",
    "grupo",
    "categoria_armazenagem",
    "vol_l_unitario",
    "vol_L_unitario",
    "venda_total",
    "nm_fabricante",
    "altura_cm",
    "peso_kg_unitario",
    "subcategoria",
    "is_pesado",
    "is_alto",
    "is_hot_zone",
    "is_nivel_alto",
    "is_nivel_inferior",
    "is_realocado",
    "location_id_atual",
    "slot_duplo",
    "produto_alocado_code",
    "grupo_alocado",
]


def normalize_code(value):
    if value is None:
        return None
    return str(value).strip().upper()


def clone_cell(source_cell, target_cell):
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)


def copy_sheet_defaults(source_ws, target_ws):
    target_ws.sheet_format.defaultColWidth = source_ws.sheet_format.defaultColWidth
    target_ws.sheet_format.defaultRowHeight = source_ws.sheet_format.defaultRowHeight
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
    target_ws.page_margins = copy(source_ws.page_margins)
    target_ws.page_setup = copy(source_ws.page_setup)
    target_ws.print_options = copy(source_ws.print_options)
    target_ws.sheet_state = source_ws.sheet_state
    for key, dim in source_ws.column_dimensions.items():
        target_dim = target_ws.column_dimensions[key]
        target_dim.width = dim.width
        target_dim.hidden = dim.hidden
        target_dim.bestFit = dim.bestFit
        target_dim.outlineLevel = dim.outlineLevel
        target_dim.collapsed = dim.collapsed


def worksheet_index(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


def clear_product_values(values):
    cleared = dict(values)
    for key in PRODUCT_COLUMNS:
        if key not in cleared:
            continue
        cleared[key] = ""
    if "product_code" in cleared:
        cleared["product_code"] = "Vazio"
    if "quantidade" in cleared:
        cleared["quantidade"] = 0
    if "is_hot_zone" in cleared:
        cleared["is_hot_zone"] = "FALSE"
    if "is_nivel_alto" in cleared:
        cleared["is_nivel_alto"] = "FALSE"
    if "is_nivel_inferior" in cleared:
        cleared["is_nivel_inferior"] = "TRUE"
    if "location_id_atual" in cleared:
        cleared["location_id_atual"] = cleared.get("location_id")
    if "produto_alocado_code" in cleared:
        cleared["produto_alocado_code"] = "Vazio"
    return cleared


def main():
    wb = openpyxl.load_workbook(INPUT_PATH)
    wb_skeleton = openpyxl.load_workbook(SKELETON_PATH, data_only=False)
    wb_join = openpyxl.load_workbook(JOIN_PATH, data_only=False)

    ws_mix = wb["Mix"]
    ws_base = wb["Base"]
    ws_plan = wb[PLAN_SHEET]
    ws_skeleton = wb_skeleton[SKELETON_SHEET]
    ws_join = wb_join[JOIN_SHEET]

    mix_idx = worksheet_index(ws_mix)
    base_idx = worksheet_index(ws_base)
    skeleton_idx = worksheet_index(ws_skeleton)
    join_idx = worksheet_index(ws_join)

    valid_codes = set()
    for r in range(2, ws_mix.max_row + 1):
        code = normalize_code(ws_mix.cell(r, mix_idx["Código"]).value)
        if code:
            valid_codes.add(code)
    base_codes = {
        normalize_code(ws_base.cell(r, base_idx["product_code"]).value)
        for r in range(2, ws_base.max_row + 1)
        if normalize_code(ws_base.cell(r, base_idx["product_code"]).value)
    }
    valid_codes &= base_codes

    join_headers = [ws_join.cell(1, c).value for c in range(1, ws_join.max_column + 1)]

    assignments_by_location = defaultdict(list)
    join_row_order = []
    for r in range(2, ws_join.max_row + 1):
        code = normalize_code(ws_join.cell(r, join_idx["product_code"]).value)
        if code not in valid_codes:
            continue
        location_id = ws_join.cell(r, join_idx["location_id"]).value
        values = {header: ws_join.cell(r, join_idx[header]).value for header in join_headers}
        assignments_by_location[location_id].append({"row": r, "values": values})
        join_row_order.append((location_id, r))

    used_assignment_rows = set()

    ws_plan.delete_rows(1, ws_plan.max_row)
    copy_sheet_defaults(ws_skeleton, ws_plan)

    for c in range(1, ws_skeleton.max_column + 1):
        clone_cell(ws_skeleton.cell(1, c), ws_plan.cell(1, c))
    ws_plan.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws_skeleton.max_column)}1"

    out_row = 2
    location_offsets = defaultdict(int)
    for source_row in range(2, ws_skeleton.max_row + 1):
        location_id = ws_skeleton.cell(source_row, skeleton_idx["location_id"]).value
        assigned_list = assignments_by_location.get(location_id, [])
        offset = location_offsets[location_id]

        if offset < len(assigned_list):
            assigned = assigned_list[offset]
            used_assignment_rows.add(assigned["row"])
            row_values = assigned["values"]
            style_source_ws = ws_join
            style_source_row = assigned["row"]
        else:
            skeleton_values = {header: ws_skeleton.cell(source_row, skeleton_idx[header]).value for header in join_headers}
            row_values = clear_product_values(skeleton_values)
            style_source_ws = ws_skeleton
            style_source_row = source_row

        location_offsets[location_id] += 1

        src_dim = style_source_ws.row_dimensions[style_source_row]
        if src_dim.height is not None:
            ws_plan.row_dimensions[out_row].height = src_dim.height
        for c, header in enumerate(join_headers, start=1):
            clone_cell(style_source_ws.cell(style_source_row, c), ws_plan.cell(out_row, c))
            ws_plan.cell(out_row, c).value = row_values.get(header)
        out_row += 1

    for location_id, source_row in join_row_order:
        if source_row in used_assignment_rows:
            continue
        src_dim = ws_join.row_dimensions[source_row]
        if src_dim.height is not None:
            ws_plan.row_dimensions[out_row].height = src_dim.height
        for c in range(1, ws_join.max_column + 1):
            clone_cell(ws_join.cell(source_row, c), ws_plan.cell(out_row, c))
        out_row += 1

    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
