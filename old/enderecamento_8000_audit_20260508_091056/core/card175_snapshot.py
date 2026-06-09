from __future__ import annotations

import io
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .gsheets_client import CREDENTIALS_DIR, GSheetsClient
from .utils import normalize_string, parse_number

CARD175_PLAN_SHEET = "Plano_Enderecamento_Final_card_175"
CARD175_BASE_SHEET = "Plano_Enderecamento_Final_card_175_base"
CARD175_INVALID_SHEET = "Card175_Enderecos_Invalidos"
CARD175_CHANGELOG_SHEET = "Log_Alteracoes_card_175"
WORKING_PLAN_SHEET = "Plano_Enderecamento_Final"
MAP_SHEET_FALLBACK = "Mapa_Final_Escaninhos"
CARD175_CONTEXT_PATH = CREDENTIALS_DIR / "card175_context.json"
UNALLOCATED_ID = "nao-alocado"
PRANCHETA_ID = "prancheta"
UNALLOCATED_LABEL = "NÃO ALOCADO"
ALPHABET = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

CARD175_REQUIRED_PLAN_HEADERS = [
    "location_id",
    "galpao_id",
    "rua_num",
    "equipamento_num",
    "tipo_equipamento",
    "nivel",
    "escaninho_num_no_nivel",
    "capacidade_l",
    "tipo_equipamento_final",
    "product_code",
    "product_name",
    "quantidade",
    "curva",
    "grupo",
    "categoria_armazenagem",
    "vol_l_unitario",
    "vol_L_unitario",
    "venda_total",
    "nm_fabricante",
    "altura_cm",
    "peso_kg_unitario",
    "subcategoria",
    "is_pesado",
    "is_alto",
    "is_hot_zone",
    "is_nivel_alto",
    "is_nivel_inferior",
    "is_realocado",
    "location_id_atual",
    "slot_duplo",
    "produto_alocado_code",
    "grupo_alocado",
]


def _save_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _load_json(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _normalize_header(value: Any) -> str:
    text = normalize_string(value).strip().lower()
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _find_index(headers: list[str], *aliases: str) -> int:
    aliases_norm = {_normalize_header(alias) for alias in aliases}
    for idx, header in enumerate(headers):
        if _normalize_header(header) in aliases_norm:
            return idx
    return -1


def _append_rows_chunked(client: GSheetsClient, sheet_name: str, rows: list[list[Any]], chunk_size: int = 400) -> None:
    if not rows:
        return
    for i in range(0, len(rows), chunk_size):
        client.append_rows(sheet_name, rows[i : i + chunk_size])


def _resolve_plan_source_sheet(client: GSheetsClient) -> tuple[str, str]:
    names = client.list_sheet_names()
    if WORKING_PLAN_SHEET in names:
        return WORKING_PLAN_SHEET, "plano"
    for name in names:
        norm = _normalize_header(name)
        if "plano_enderecamento" in norm:
            return name, "plano"
    if MAP_SHEET_FALLBACK in names:
        return MAP_SHEET_FALLBACK, "mapa"
    for name in names:
        norm = _normalize_header(name)
        if "mapa_final_escaninhos" in norm:
            return name, "mapa"
    raise ValueError(
        "Não encontrei aba de plano de endereçamento nem Mapa_Final_Escaninhos na planilha informada."
    )


def _ensure_required_headers(headers: list[str]) -> list[str]:
    normalized = {_normalize_header(item) for item in headers if normalize_string(item)}
    out = list(headers)
    for required in CARD175_REQUIRED_PLAN_HEADERS:
        key = _normalize_header(required)
        if key and key not in normalized:
            out.append(required)
            normalized.add(key)
    return out


def _to_float_qty(value: Any) -> float:
    parsed = parse_number(value)
    if parsed is None:
        return 0.0
    return float(parsed)


def _to_int_if_whole(value: float) -> int | float:
    if abs(value - round(value)) < 1e-9:
        return int(round(value))
    return round(value, 4)


def _string_variants(value: Any, include_numeric: bool = True) -> set[str]:
    raw = normalize_string(value).upper()
    out: set[str] = set()
    if raw:
        out.add(raw)
        out.add(raw.lstrip("0") or "0")
        if raw.startswith("R") and raw[1:].isdigit():
            out.add(raw[1:].lstrip("0") or "0")
    if include_numeric:
        numeric_like = bool(re.fullmatch(r"[A-Z]?\d+", raw))
        num = parse_number(value)
        if num is not None and (numeric_like or raw.isdigit()):
            int_num = int(round(num))
            out.add(str(int_num))
            out.add(f"{int_num:03d}")
            out.add(f"R{int_num}")
    return {item for item in out if item}


def _register_addr_mapping(
    addr_map: dict[str, str],
    location_id: str,
    galpao: Any,
    rua: Any,
    posicao: Any,
    escaninho: Any,
) -> None:
    loc = normalize_string(location_id).strip()
    gal = normalize_string(galpao).upper().strip()
    if not loc or not gal:
        return
    ruas = _string_variants(rua, include_numeric=True) or {normalize_string(rua).upper()}
    posicoes = _string_variants(posicao, include_numeric=True) or {normalize_string(posicao).upper()}
    escaninhos = _string_variants(escaninho, include_numeric=False) or {normalize_string(escaninho).upper()}
    for rua_val in ruas:
        for pos_val in posicoes:
            for esc_val in escaninhos:
                key = f"{gal}|{rua_val}|{pos_val}|{esc_val}"
                addr_map.setdefault(key, loc)


def _extract_location_parts(location_id: str) -> tuple[str, str, str, str] | None:
    text = normalize_string(location_id).upper()
    match = re.match(r"^([A-Z0-9_]+)-R(\d+)-(?:E)?(\d+)-(.+)$", text)
    if not match:
        return None
    galpao = match.group(1)
    rua = match.group(2)
    pos = match.group(3)
    esc = match.group(4)
    return galpao, rua, pos, esc


def _load_card175_upload_rows(file_bytes: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb[wb.sheetnames[0]]
    values = list(sheet.iter_rows(min_row=1, values_only=True))
    if not values:
        return []
    headers = [normalize_string(col) for col in values[0]]
    header_norm = [_normalize_header(h) for h in headers]

    idx_id_local = _find_index(header_norm, "id_localizacao")
    idx_galpao = _find_index(header_norm, "galpao")
    idx_rua = _find_index(header_norm, "rua")
    idx_pos = _find_index(header_norm, "posicao_pallete", "posicao_pallet", "posicao", "estante")
    idx_esc = _find_index(header_norm, "escaninho_nivel", "escaninho", "nivel", "posicao_nivel")
    idx_code = _find_index(header_norm, "cod_produto", "codigo_produto", "sku", "product_code")
    idx_desc = _find_index(header_norm, "desc_produto", "descricao_produto", "product_name", "descricao")
    idx_qty = _find_index(header_norm, "quantidade", "qtd", "estoque")

    if idx_code == -1:
        raise ValueError("Arquivo do card não possui coluna de código do produto (cod_produto).")

    aggregated: dict[tuple[str, str, str, str, str, str, str], dict[str, Any]] = {}
    for raw in values[1:]:
        row = list(raw)
        code = normalize_string(row[idx_code] if idx_code < len(row) else "")
        if not code:
            continue
        id_local = normalize_string(row[idx_id_local] if idx_id_local >= 0 and idx_id_local < len(row) else "")
        galpao = normalize_string(row[idx_galpao] if idx_galpao >= 0 and idx_galpao < len(row) else "")
        rua = normalize_string(row[idx_rua] if idx_rua >= 0 and idx_rua < len(row) else "")
        pos = normalize_string(row[idx_pos] if idx_pos >= 0 and idx_pos < len(row) else "")
        esc = normalize_string(row[idx_esc] if idx_esc >= 0 and idx_esc < len(row) else "")
        desc = normalize_string(row[idx_desc] if idx_desc >= 0 and idx_desc < len(row) else "")
        qty = _to_float_qty(row[idx_qty] if idx_qty >= 0 and idx_qty < len(row) else 0)
        key = (id_local, galpao, rua, pos, esc, code, desc)
        if key not in aggregated:
            aggregated[key] = {
                "id_localizacao": id_local,
                "galpao": galpao,
                "rua": rua,
                "posicao_pallete": pos,
                "escaninho_nivel": esc,
                "cod_produto": code,
                "desc_produto": desc,
                "quantidade": 0.0,
            }
        aggregated[key]["quantidade"] += qty
    return list(aggregated.values())


def _build_external_location_maps(client: GSheetsClient) -> tuple[dict[str, str], dict[str, str]]:
    id_map: dict[str, str] = {}
    addr_map: dict[str, str] = {}
    for sheet_name in client.list_sheet_names():
        values = client.read_values(sheet_name)
        if not values or len(values) < 2:
            continue
        headers = [normalize_string(h) for h in values[0]]
        header_norm = [_normalize_header(h) for h in headers]
        idx_loc = _find_index(header_norm, "location_id", "location")
        if idx_loc == -1:
            continue
        idx_id_local = _find_index(header_norm, "id_localizacao")
        idx_galpao = _find_index(header_norm, "galpao", "galpao_id")
        idx_rua = _find_index(header_norm, "rua", "rua_num")
        idx_pos = _find_index(
            header_norm,
            "posicao_pallete",
            "posicao_pallet",
            "posicao",
            "equipamento_num",
            "equipamento",
            "estante",
        )
        idx_esc_composto = _find_index(header_norm, "escaninho_nivel")
        idx_nivel = _find_index(header_norm, "nivel")
        idx_esc_no_nivel = _find_index(header_norm, "escaninho_num_no_nivel", "escaninho")

        for raw in values[1:]:
            row = list(raw)
            location_id = normalize_string(row[idx_loc] if idx_loc < len(row) else "")
            if not location_id:
                continue
            if idx_id_local >= 0 and idx_id_local < len(row):
                id_local = normalize_string(row[idx_id_local])
                if id_local:
                    id_map.setdefault(id_local, location_id)
            if idx_galpao >= 0 and idx_rua >= 0 and idx_pos >= 0:
                gal = row[idx_galpao] if idx_galpao < len(row) else ""
                rua = row[idx_rua] if idx_rua < len(row) else ""
                pos = row[idx_pos] if idx_pos < len(row) else ""
                esc = ""
                if idx_esc_composto >= 0 and idx_esc_composto < len(row):
                    esc = normalize_string(row[idx_esc_composto])
                if not esc and idx_nivel >= 0 and idx_nivel < len(row) and idx_esc_no_nivel >= 0 and idx_esc_no_nivel < len(row):
                    nivel = normalize_string(row[idx_nivel])
                    esc_no_nivel = normalize_string(row[idx_esc_no_nivel])
                    if nivel and esc_no_nivel:
                        esc = f"{nivel}{esc_no_nivel}"
                if not esc and idx_esc_no_nivel >= 0 and idx_esc_no_nivel < len(row):
                    esc = normalize_string(row[idx_esc_no_nivel])
                if not esc and idx_nivel >= 0 and idx_nivel < len(row):
                    esc = normalize_string(row[idx_nivel])
                _register_addr_mapping(addr_map, location_id, gal, rua, pos, esc)
    return id_map, addr_map


def _build_base_products_map(client: GSheetsClient) -> dict[str, dict[str, Any]]:
    try:
        rows = client.read_sheet("Base_Produtos")
    except Exception:
        return {}
    mapping: dict[str, dict[str, Any]] = {}
    for row in rows:
        row_norm = {_normalize_header(k): v for k, v in row.items()}
        code = normalize_string(
            row_norm.get("product_code")
            or row_norm.get("codigo")
            or row_norm.get("cod_produto")
            or row_norm.get("sku")
        )
        if not code:
            continue
        mapping[code] = row_norm
    return mapping


def _set_card175_context(payload: dict[str, Any]) -> None:
    _save_json(CARD175_CONTEXT_PATH, payload)


def get_card175_context(sheet_id: str | None = None) -> dict[str, Any] | None:
    payload = _load_json(CARD175_CONTEXT_PATH)
    if not isinstance(payload, dict):
        return None
    if not payload.get("enabled"):
        return None
    if sheet_id and normalize_string(payload.get("sheet_id")) != normalize_string(sheet_id):
        return None
    return payload


def append_card175_change_logs(sheet_id: str, moves: list[dict[str, Any]], user: str = "local") -> dict[str, Any]:
    context = get_card175_context(sheet_id)
    if not context:
        return {"success": True, "logged": 0}

    changes_sheet = normalize_string(context.get("changes_sheet")) or CARD175_CHANGELOG_SHEET
    source_file = normalize_string(context.get("source_file"))
    client = GSheetsClient(sheet_id)
    client.ensure_sheet(changes_sheet)

    headers = [
        "timestamp",
        "usuario",
        "product_code",
        "location_id_original",
        "location_id_atual",
        "origem_operacao",
        "arquivo_card_175",
    ]
    existing = client.read_values(changes_sheet)
    if not existing or [normalize_string(h) for h in existing[0]] != headers:
        client.clear_sheet(changes_sheet)
        client.append_rows(changes_sheet, [headers])

    now = datetime.now()
    timestamp = now.strftime("%Y-%m-%d %H:%M:%S")
    rows_to_append: list[list[Any]] = []
    for move in moves:
        loc_old_raw = normalize_string(move.get("locAnteriorId"))
        loc_new_raw = normalize_string(move.get("locNovoId"))
        if not loc_new_raw or loc_new_raw == PRANCHETA_ID:
            continue
        loc_old = loc_old_raw.replace("bin-", "") if loc_old_raw else UNALLOCATED_LABEL
        loc_new = loc_new_raw.replace("bin-", "") if loc_new_raw else UNALLOCATED_LABEL
        if loc_old_raw == UNALLOCATED_ID:
            loc_old = UNALLOCATED_LABEL
        if loc_new_raw == UNALLOCATED_ID:
            loc_new = UNALLOCATED_LABEL
        if normalize_string(loc_old) == normalize_string(loc_new):
            continue
        rows_to_append.append(
            [
                timestamp,
                user,
                normalize_string(move.get("productCode")),
                loc_old,
                loc_new,
                "DASHBOARD",
                source_file,
            ]
        )
    if rows_to_append:
        _append_rows_chunked(client, changes_sheet, rows_to_append)
    return {"success": True, "logged": len(rows_to_append), "changes_sheet": changes_sheet}


def import_card175_snapshot(
    sheet_id: str,
    file_bytes: bytes,
    source_file_name: str,
    user: str = "local",
) -> dict[str, Any]:
    if not file_bytes:
        return {"success": False, "error": "Arquivo vazio."}

    client = GSheetsClient(sheet_id)
    source_plan_sheet, source_kind = _resolve_plan_source_sheet(client)
    source_values = client.read_values(source_plan_sheet)
    if not source_values or len(source_values) < 2:
        return {"success": False, "error": f"Aba {source_plan_sheet} vazia."}

    source_headers = [normalize_string(h) for h in source_values[0]]
    plan_headers = _ensure_required_headers(source_headers)
    if not any(plan_headers):
        return {"success": False, "error": f"Cabeçalho inválido em {source_plan_sheet}."}

    idx_loc = _find_index(plan_headers, "location_id")
    idx_code = _find_index(plan_headers, "product_code")
    if idx_code == -1:
        idx_code = _find_index(plan_headers, "produto_alocado_code")
    if idx_loc == -1 or idx_code == -1:
        return {"success": False, "error": "Plano de origem sem colunas location_id/product_code."}

    idx_slot_duplo = _find_index(plan_headers, "slot_duplo")
    idx_name = _find_index(plan_headers, "product_name", "desc_produto")
    idx_qty = _find_index(plan_headers, "quantidade")
    idx_loc_atual = _find_index(plan_headers, "location_id_atual")
    idx_prod_alocado = _find_index(plan_headers, "produto_alocado_code")
    if idx_prod_alocado == -1:
        idx_prod_alocado = _find_index(plan_headers, "product_code")
    idx_group_alocado = _find_index(plan_headers, "grupo_alocado")

    extra_cols = len(plan_headers) - len(source_headers)
    data_rows = []
    for row in source_values[1:]:
        row_list = list(row)
        padded = row_list + [None] * (len(source_headers) - len(row_list))
        if extra_cols > 0:
            padded.extend([""] * extra_cols)
        data_rows.append(padded)
    template_by_loc: dict[str, list[Any]] = {}
    ordered_locations: list[str] = []
    for row in data_rows:
        loc = normalize_string(row[idx_loc]) if idx_loc < len(row) else ""
        if not loc:
            continue
        if loc not in template_by_loc:
            template_by_loc[loc] = list(row[: len(plan_headers)])
            ordered_locations.append(loc)

    if not template_by_loc:
        return {"success": False, "error": "Nenhum location_id válido no plano de origem."}

    upload_rows = _load_card175_upload_rows(file_bytes)
    if not upload_rows:
        return {"success": False, "error": "Arquivo do card sem linhas válidas."}

    id_map, addr_map = _build_external_location_maps(client)
    for loc in ordered_locations:
        parts = _extract_location_parts(loc)
        if not parts:
            continue
        gal, rua, pos, esc = parts
        _register_addr_mapping(addr_map, loc, gal, rua, pos, esc)

    aggregated_by_loc: dict[str, list[dict[str, Any]]] = defaultdict(list)
    unresolved_groups: dict[str, dict[str, Any]] = {}
    invalid_rows: list[list[Any]] = []
    unresolved = 0
    for item in upload_rows:
        id_local = normalize_string(item.get("id_localizacao"))
        code = normalize_string(item.get("cod_produto"))
        if not code:
            continue
        resolved_loc = ""
        if id_local and id_local in id_map:
            resolved_loc = id_map[id_local]
        if not resolved_loc:
            gal = normalize_string(item.get("galpao")).upper()
            rua_variants = _string_variants(item.get("rua"), include_numeric=True) or {normalize_string(item.get("rua")).upper()}
            pos_variants = _string_variants(item.get("posicao_pallete"), include_numeric=True) or {normalize_string(item.get("posicao_pallete")).upper()}
            esc_variants = _string_variants(item.get("escaninho_nivel"), include_numeric=False) or {normalize_string(item.get("escaninho_nivel")).upper()}
            for rua in rua_variants:
                if resolved_loc:
                    break
                for pos in pos_variants:
                    if resolved_loc:
                        break
                    for esc in esc_variants:
                        key = f"{gal}|{rua}|{pos}|{esc}"
                        mapped = addr_map.get(key)
                        if mapped:
                            resolved_loc = mapped
                            break
        if not resolved_loc or resolved_loc not in template_by_loc:
            unresolved += 1
            gal_raw = normalize_string(item.get("galpao")).upper()
            rua_raw = normalize_string(item.get("rua")).upper()
            pos_raw = normalize_string(item.get("posicao_pallete")).upper()
            esc_raw = normalize_string(item.get("escaninho_nivel")).upper()
            group_key = f"{gal_raw}|{rua_raw}|{pos_raw}|{esc_raw}"
            group = unresolved_groups.get(group_key)
            if not group:
                group = {
                    "galpao": gal_raw or "SEM_GALPAO",
                    "rua": rua_raw,
                    "posicao": pos_raw,
                    "escaninho": esc_raw,
                    "id_localizacao": normalize_string(item.get("id_localizacao")),
                    "items": [],
                }
                unresolved_groups[group_key] = group
            group["items"].append(item)
            continue
        aggregated_by_loc[resolved_loc].append(item)

    base_products = _build_base_products_map(client)
    plan_rows_out: list[list[Any]] = []
    location_slot_counts: dict[str, int] = {}
    product_columns = [
        "product_code",
        "produto_alocado_code",
        "product_name",
        "quantidade",
        "curva",
        "grupo",
        "grupo_alocado",
        "categoria_armazenagem",
        "vol_l_unitario",
        "vol_l_unitario",
        "venda_total",
        "nm_fabricante",
        "altura_cm",
        "peso_kg_unitario",
        "subcategoria",
        "is_pesado",
        "is_alto",
        "is_pequeno",
        "is_fragil",
        "degelo",
        "metodo",
        "location_id_atual",
    ]
    product_col_indices = [idx for idx, h in enumerate(plan_headers) if _normalize_header(h) in product_columns]

    def clear_product_columns(row_data: list[Any]) -> None:
        for idx in product_col_indices:
            if idx < len(row_data):
                row_data[idx] = ""

    def apply_product(row_data: list[Any], location_id: str, product_data: dict[str, Any]) -> None:
        clear_product_columns(row_data)
        code = normalize_string(product_data.get("cod_produto"))
        desc = normalize_string(product_data.get("desc_produto"))
        qty = _to_int_if_whole(_to_float_qty(product_data.get("quantidade")))
        base = base_products.get(code, {})
        if idx_code >= 0:
            row_data[idx_code] = code
        if idx_prod_alocado >= 0:
            row_data[idx_prod_alocado] = code
        if idx_name >= 0:
            row_data[idx_name] = normalize_string(base.get("product_name") or desc)
        if idx_qty >= 0:
            row_data[idx_qty] = qty
        if idx_loc_atual >= 0:
            row_data[idx_loc_atual] = location_id
        if idx_group_alocado >= 0:
            row_data[idx_group_alocado] = normalize_string(base.get("grupo"))

        for idx, header in enumerate(plan_headers):
            key = _normalize_header(header)
            if key in {"product_code", "produto_alocado_code", "product_name", "quantidade", "location_id_atual", "grupo_alocado"}:
                continue
            if key in base and normalize_string(base.get(key)) != "":
                row_data[idx] = base.get(key)

    def apply_empty(row_data: list[Any], location_id: str) -> None:
        clear_product_columns(row_data)
        if idx_code >= 0:
            row_data[idx_code] = "Vazio"
        if idx_prod_alocado >= 0:
            row_data[idx_prod_alocado] = "Vazio"
        if idx_name >= 0:
            row_data[idx_name] = ""
        if idx_qty >= 0:
            row_data[idx_qty] = 0
        if idx_loc_atual >= 0:
            row_data[idx_loc_atual] = location_id

    overflow_count = 0
    for location_id in ordered_locations:
        template = template_by_loc[location_id]
        rows_here = sorted(
            aggregated_by_loc.get(location_id, []),
            key=lambda item: _to_float_qty(item.get("quantidade")),
            reverse=True,
        )
        if not rows_here:
            out_row = list(template)
            apply_empty(out_row, location_id)
            plan_rows_out.append(out_row)
            location_slot_counts[location_id] = 0
            continue

        keep = rows_here[:2]
        overflow = rows_here[2:]
        if overflow:
            overflow_count += len(overflow)
            for item in overflow:
                invalid_rows.append(
                    [
                        "MAIS_DE_2_PRODUTOS_NO_ESCANINHO",
                        normalize_string(item.get("id_localizacao")),
                        normalize_string(item.get("galpao")),
                        normalize_string(item.get("rua")),
                        normalize_string(item.get("posicao_pallete")),
                        normalize_string(item.get("escaninho_nivel")),
                        normalize_string(item.get("cod_produto")),
                        normalize_string(item.get("desc_produto")),
                        _to_int_if_whole(_to_float_qty(item.get("quantidade"))),
                        location_id,
                    ]
                )

        for item in keep:
            out_row = list(template)
            apply_product(out_row, location_id, item)
            plan_rows_out.append(out_row)
        location_slot_counts[location_id] = len(keep)

    def _safe_int(value: Any) -> int | None:
        parsed = parse_number(value)
        if parsed is None:
            return None
        try:
            return int(round(float(parsed)))
        except Exception:
            return None

    idx_galpao_id = _find_index(plan_headers, "galpao_id", "galpao")
    idx_rua_num = _find_index(plan_headers, "rua_num", "rua")
    idx_equip_num = _find_index(plan_headers, "equipamento_num", "equipamento")
    idx_tipo_equip = _find_index(plan_headers, "tipo_equipamento")
    idx_tipo_equip_final = _find_index(plan_headers, "tipo_equipamento_final")
    idx_nivel = _find_index(plan_headers, "nivel")
    idx_esc_no_nivel = _find_index(plan_headers, "escaninho_num_no_nivel", "escaninho")
    idx_capacidade = _find_index(plan_headers, "capacidade_l")
    idx_exclusivo = _find_index(plan_headers, "exclusivo_para")

    virtual_rua_num = 99
    next_virtual_equip = 900
    used_virtual_keys: set[tuple[int, int]] = set()
    virtual_locations_count = 0

    for group_key in sorted(unresolved_groups.keys()):
        group = unresolved_groups[group_key]
        group_items = sorted(
            group.get("items", []),
            key=lambda item: _to_float_qty(item.get("quantidade")),
            reverse=True,
        )
        if not group_items:
            continue

        equip_num_candidate = _safe_int(group.get("posicao"))
        if equip_num_candidate is None or equip_num_candidate <= 0:
            equip_num_candidate = next_virtual_equip

        while (virtual_rua_num, equip_num_candidate) in used_virtual_keys:
            equip_num_candidate += 1
        used_virtual_keys.add((virtual_rua_num, equip_num_candidate))
        next_virtual_equip = max(next_virtual_equip, equip_num_candidate + 1)

        source_addr = f"{group.get('galpao')}|{group.get('rua')}|{group.get('posicao')}|{group.get('escaninho')}"

        pair_index = 0
        for start in range(0, len(group_items), 2):
            pair = group_items[start : start + 2]
            if not pair:
                continue

            level_num = pair_index // 7 + 1
            esc_idx = pair_index % 7 + 1
            esc_label = ALPHABET[esc_idx - 1]
            esc_text = f"{level_num}{esc_label}"
            loc_id = f"{group.get('galpao') or 'SEM_GALPAO'}-R{virtual_rua_num}-{equip_num_candidate:03d}-{esc_text}"

            base_row = [""] * len(plan_headers)
            if idx_loc >= 0:
                base_row[idx_loc] = loc_id
            if idx_galpao_id >= 0:
                base_row[idx_galpao_id] = group.get("galpao") or "SEM_GALPAO"
            if idx_rua_num >= 0:
                base_row[idx_rua_num] = virtual_rua_num
            if idx_equip_num >= 0:
                base_row[idx_equip_num] = equip_num_candidate
            if idx_tipo_equip >= 0:
                base_row[idx_tipo_equip] = "prateleira"
            if idx_tipo_equip_final >= 0:
                base_row[idx_tipo_equip_final] = "prateleira"
            if idx_nivel >= 0:
                base_row[idx_nivel] = level_num
            if idx_esc_no_nivel >= 0:
                base_row[idx_esc_no_nivel] = esc_idx
            if idx_capacidade >= 0:
                base_row[idx_capacidade] = 8
            if idx_exclusivo >= 0:
                base_row[idx_exclusivo] = f"CARD175_FORA_MAPA:{source_addr}"

            for item in pair:
                out_row = list(base_row)
                apply_product(out_row, loc_id, item)
                plan_rows_out.append(out_row)
                invalid_rows.append(
                    [
                        "ENDERECO_NAO_MAPEADO_REPRESENTADO",
                        normalize_string(item.get("id_localizacao")) or normalize_string(group.get("id_localizacao")),
                        normalize_string(item.get("galpao")) or normalize_string(group.get("galpao")),
                        normalize_string(item.get("rua")) or normalize_string(group.get("rua")),
                        normalize_string(item.get("posicao_pallete")) or normalize_string(group.get("posicao")),
                        normalize_string(item.get("escaninho_nivel")) or normalize_string(group.get("escaninho")),
                        normalize_string(item.get("cod_produto")),
                        normalize_string(item.get("desc_produto")),
                        _to_int_if_whole(_to_float_qty(item.get("quantidade"))),
                        loc_id,
                    ]
                )
            location_slot_counts[loc_id] = len(pair)
            virtual_locations_count += 1
            pair_index += 1

    if idx_slot_duplo == -1:
        plan_headers.append("slot_duplo")
        idx_slot_duplo = len(plan_headers) - 1
        for i in range(len(plan_rows_out)):
            plan_rows_out[i] = list(plan_rows_out[i]) + [""]

    for row in plan_rows_out:
        loc = normalize_string(row[idx_loc]) if idx_loc < len(row) else ""
        slot_count = location_slot_counts.get(loc, 0)
        row[idx_slot_duplo] = "SIM" if slot_count >= 2 else "NAO"

    client.clear_sheet(CARD175_PLAN_SHEET)
    _append_rows_chunked(client, CARD175_PLAN_SHEET, [plan_headers] + plan_rows_out)

    client.clear_sheet(WORKING_PLAN_SHEET)
    _append_rows_chunked(client, WORKING_PLAN_SHEET, [plan_headers] + plan_rows_out)

    client.clear_sheet(CARD175_BASE_SHEET)
    _append_rows_chunked(client, CARD175_BASE_SHEET, [plan_headers] + plan_rows_out)

    invalid_headers = [
        "motivo",
        "id_localizacao",
        "galpao",
        "rua",
        "posicao_pallete",
        "escaninho_nivel",
        "cod_produto",
        "desc_produto",
        "quantidade",
        "location_id_resolvido",
    ]
    client.clear_sheet(CARD175_INVALID_SHEET)
    _append_rows_chunked(client, CARD175_INVALID_SHEET, [invalid_headers] + invalid_rows)

    changelog_headers = [
        "timestamp",
        "usuario",
        "product_code",
        "location_id_original",
        "location_id_atual",
        "origem_operacao",
        "arquivo_card_175",
    ]
    client.clear_sheet(CARD175_CHANGELOG_SHEET)
    _append_rows_chunked(client, CARD175_CHANGELOG_SHEET, [changelog_headers])

    _set_card175_context(
        {
            "enabled": True,
            "sheet_id": sheet_id,
            "plan_sheet": CARD175_PLAN_SHEET,
            "working_sheet": WORKING_PLAN_SHEET,
            "base_sheet": CARD175_BASE_SHEET,
            "invalid_sheet": CARD175_INVALID_SHEET,
            "changes_sheet": CARD175_CHANGELOG_SHEET,
            "source_file": normalize_string(source_file_name),
            "imported_at": datetime.now().isoformat(timespec="seconds"),
        }
    )

    return {
        "success": True,
        "source_plan_sheet": source_plan_sheet,
        "source_sheet_kind": source_kind,
        "generated_sheet": CARD175_PLAN_SHEET,
        "rows_written": len(plan_rows_out),
        "locations_total": len(ordered_locations),
        "locations_with_data": sum(1 for value in location_slot_counts.values() if value > 0),
        "invalid_rows": len(invalid_rows),
        "unmapped_rows": unresolved,
        "overflow_rows": overflow_count,
        "virtual_locations": virtual_locations_count,
        "virtual_address_groups": len(unresolved_groups),
        "generated_sheet_url": client.get_sheet_url(CARD175_PLAN_SHEET),
        "working_sheet_url": client.get_sheet_url(WORKING_PLAN_SHEET),
        "invalid_sheet_url": client.get_sheet_url(CARD175_INVALID_SHEET),
        "changes_sheet_url": client.get_sheet_url(CARD175_CHANGELOG_SHEET),
    }
