#!/usr/bin/env python3
# Gera:
# - "Ainda_falta_calculado" com enderecos pendentes (local, sem Apps Script)
# - "Baixa_enderecos_antigos" com baixas (quantidade negativa) de enderecos antigos
# - "Prova_baixa_antigos" com validacao da baixa (endereco antigo + quantidade)
# Regras:
# - Exclui enderecos ja presentes na aba Endereçamento (LocalCadastrado/LocalBipado), inclusive ranges "ate"
# - Exclui enderecos bloqueados (is_blocked=1 ou is_gondola=0)
# - Mantem apenas tipos prateleira / prateleira lateral / prateleira alta

import sys
import re
import unicodedata
from datetime import datetime
from typing import Any, Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import load_workbook

DEFAULT_FILE = "planilhas/atuais/Conferência Produtos - [Pinheiros] (23).xlsx"
SHEET_LOCS = "Localizações produtos"
SHEET_END = "Endereçamento"
SHEET_BLOCK = "Enderecos_Bloqueados"
SHEET_COD = "Codigos de barras"
SHEET_STOCK = "Estoque_Atual"
SHEET_LAYOUT_PRIMARY = "Plano_Enderecamento_Final_Layout_Atual"
OUT_SHEET = "Ainda_falta_calculado"
OUT_SHEET_BAIXA = "Baixa_enderecos_antigos"
OUT_SHEET_PROVA = "Prova_baixa_antigos"
BAIXA_COLUMNS = [
    "cod_produto",
    "galpao",
    "rua",
    "estante",
    "escaninho",
    "tipo",
    "quantidade",
    "justificativa",
    "data_validade",
    "destino",
    "unidade_medida",
]
PROVA_COLUMNS = [
    "cod_produto",
    "endereco_antigo",
    "enderecos_novos_layout",
    "quantidade_estoque_antigo",
    "quantidade_baixa_gerada",
    "data_validade",
    "confere_quantidade",
    "confere_endereco_antigo",
]


def norm_simple(value: str) -> str:
    s = "" if value is None else str(value)
    s = s.strip().lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"\s+", "", s)
    return s


def _split_iferror_args(inner: str):
    depth = 0
    for i, ch in enumerate(inner):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif ch == "," and depth == 0:
            return inner[:i], inner[i + 1 :]
    return inner, ""


def cell_value(cell):
    v = cell.value
    if isinstance(v, str) and v.startswith("=IFERROR("):
        inner = v[len("=IFERROR(") :]
        if inner.endswith(")"):
            inner = inner[:-1]
        _, fallback = _split_iferror_args(inner)
        fallback = fallback.strip()
        if fallback.startswith("\"") and fallback.endswith("\""):
            fallback = fallback[1:-1].replace('""', '"')
        return fallback
    return v


def sheet_to_df(wb, sheet_name: str) -> pd.DataFrame:
    sh = wb[sheet_name]
    rows = []
    for r in sh.iter_rows(min_row=1, max_row=sh.max_row, max_col=sh.max_column):
        row = [cell_value(c) for c in r]
        rows.append(row)
    # remove trailing empty rows
    cleaned = []
    for row in rows:
        if any(v is not None and str(v).strip() != "" for v in row):
            cleaned.append(row)
    if not cleaned:
        return pd.DataFrame()
    header = ["" if h is None else str(h).strip() for h in cleaned[0]]
    data = cleaned[1:]
    df = pd.DataFrame(data, columns=header)
    return df


def normalize_addr(value: str) -> str:
    s = "" if value is None else str(value)
    s = s.strip().upper()
    if not s:
        return ""
    s = s.replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"-+", "-", s)
    parts = [p for p in s.split("-") if p]
    if len(parts) >= 3:
        seg = parts[2]
        m = re.match(r"^(\d+)([A-Z]*)$", seg)
        if m:
            parts[2] = m.group(1).zfill(3) + (m.group(2) or "")
    return "-".join(parts)


def expand_range(raw_a: str, raw_b: str) -> List[str]:
    a = normalize_addr(raw_a)
    b = normalize_addr(raw_b)
    if not a and not b:
        return []
    if not a:
        return [b] if b else []
    if not b:
        return [a]
    if a == b:
        return [a]

    pa = a.split("-")

    # Reconstrui o final quando vier curto (ex: "3G", "003-3G", "R1-003-3G")
    raw_b_str = "" if raw_b is None else str(raw_b).strip()
    if raw_b_str and "-" not in raw_b_str and len(pa) >= 3:
        b = normalize_addr(f"{pa[0]}-{pa[1]}-{pa[2]}-{raw_b_str}")
    else:
        pb_tmp = normalize_addr(raw_b_str).split("-") if raw_b_str else []
        if raw_b_str and len(pb_tmp) < len(pa) and len(pa) >= 3:
            parts = re.split(r"[-–—]", raw_b_str)
            parts = [p for p in parts if p]
            if len(parts) == 1:
                b = normalize_addr(f"{pa[0]}-{pa[1]}-{pa[2]}-{parts[0]}")
            elif len(parts) == 2:
                b = normalize_addr(f"{pa[0]}-{pa[1]}-{parts[0]}-{parts[1]}")
            elif len(parts) == 3:
                b = normalize_addr(f"{pa[0]}-{parts[0]}-{parts[1]}-{parts[2]}")

    pb = b.split("-")
    if len(pa) < 4 or len(pb) < 4:
        return [a, b]
    if pa[0] != pb[0] or pa[1] != pb[1] or pa[2] != pb[2]:
        return [a, b]

    la = pa[3]
    lb = pb[3]
    ma = re.match(r"^(\d+)([A-Z])?$", la)
    mb = re.match(r"^(\d+)([A-Z])?$", lb)
    if not ma or not mb:
        return [a, b]
    if ma.group(1) != mb.group(1):
        return [a, b]

    num = ma.group(1)
    sa = ma.group(2) or ""
    sb = mb.group(2) or ""

    if sa and sb:
        start = ord(sa)
        end = ord(sb)
        if start > end:
            start, end = end, start
        return [f"{pa[0]}-{pa[1]}-{pa[2]}-{num}{chr(c)}" for c in range(start, end + 1)]

    if not sa and not sb:
        try:
            ia = int(la)
            ib = int(lb)
        except ValueError:
            return [a, b]
        start = min(ia, ib)
        end = max(ia, ib)
        return [f"{pa[0]}-{pa[1]}-{pa[2]}-{i}" for i in range(start, end + 1)]

    return [a, b]


def parse_endereco_cell(value: str) -> List[str]:
    if value is None:
        return []
    raw = str(value).strip()
    if not raw:
        return []
    parts = re.split(r"[,;\n|]", raw)
    out = []
    for p in parts:
        p = str(p).strip()
        if not p:
            continue
        if re.search(r"\bat[eé]\b", p, flags=re.IGNORECASE):
            ab = re.split(r"\bat[eé]\b", p, flags=re.IGNORECASE, maxsplit=1)
            a = (ab[0] if len(ab) > 0 else "").strip()
            b = (ab[1] if len(ab) > 1 else "").strip()
            if a and b:
                out.extend(expand_range(a, b))
            else:
                n = normalize_addr(p)
                if n:
                    out.append(n)
        else:
            n = normalize_addr(p)
            if n:
                out.append(n)
    return out


def find_col(cols, *cands):
    norm = [norm_simple(c) for c in cols]
    for cand in cands:
        c = norm_simple(cand)
        if c in norm:
            return norm.index(c)
    return -1


def parse_flag(value):
    if value is True:
        return 1
    if value is False:
        return 0
    raw = "" if value is None else str(value).strip()
    if raw == "":
        return None
    v = norm_simple(raw)
    if v in ("1", "true", "sim", "yes", "y"):
        return 1
    if v in ("0", "false", "nao", "n"):
        return 0
    try:
        n = float(raw.replace(",", "."))
        return 0 if n == 0 else 1
    except Exception:
        return None


def parse_number(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if pd.isna(value):
            return None
        return float(value)
    raw = str(value).strip()
    if not raw:
        return None
    raw = raw.replace(" ", "")
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        return float(raw)
    except Exception:
        return None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_valid_address(value: str) -> bool:
    addr = normalize_addr(value)
    if not addr:
        return False
    parts = [p for p in addr.split("-") if p]
    return len(parts) >= 4


def normalize_estante(value: Any) -> str:
    s = clean_text(value).upper()
    if not s:
        return ""
    n = parse_number(s)
    if n is not None and abs(n - round(n)) < 1e-9:
        return str(int(round(n))).zfill(3)
    m = re.match(r"^(\d+)([A-Z]*)$", s)
    if m:
        return m.group(1).zfill(3) + (m.group(2) or "")
    return s


def split_address_parts(value: str) -> Optional[Tuple[str, str, str, str]]:
    addr = normalize_addr(value)
    if not is_valid_address(addr):
        return None
    parts = [p for p in addr.split("-") if p]
    galpao = parts[0]
    rua = parts[1]
    estante = normalize_estante(parts[2])
    escaninho = "-".join(parts[3:])
    if not galpao or not rua or not estante or not escaninho:
        return None
    return galpao, rua, estante, escaninho


def normalize_validade(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        if pd.isna(value):
            return ""
        return value.to_pydatetime().date()
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "date") and not isinstance(value, str):
        try:
            return value.date()
        except Exception:
            pass
    return clean_text(value)


def to_compact_number(value: float) -> Any:
    rounded_int = int(round(value))
    if abs(value - rounded_int) < 1e-9:
        return rounded_int
    return round(value, 6)


def load_layout_df(wb) -> Tuple[pd.DataFrame, str]:
    candidates = [SHEET_LAYOUT_PRIMARY, SHEET_LOCS]
    for name in candidates:
        if name in wb.sheetnames:
            df = sheet_to_df(wb, name)
            if not df.empty:
                return df, name
    raise RuntimeError(
        f"Nenhuma aba de layout encontrada entre: {', '.join(candidates)}"
    )


def collect_new_addresses_by_product(df_layout: pd.DataFrame) -> Dict[str, Set[str]]:
    cols = list(df_layout.columns)
    idx_cod = find_col(cols, "product_code", "cod_produto", "produtoid", "codigo")
    idx_loc = find_col(cols, "location_id", "endereco", "endereco_generated")
    idx_galpao = find_col(cols, "galpao", "galpao_id")
    idx_rua = find_col(cols, "rua", "rua_num")
    idx_est = find_col(cols, "equipamento_num", "posicao_pallete", "posicao", "estante")
    idx_esc = find_col(cols, "escaninho_nivel", "escaninho")
    idx_nivel = find_col(cols, "nivel")
    idx_esc_num = find_col(cols, "escaninho_num_no_nivel")

    if idx_cod < 0:
        raise RuntimeError("Coluna product_code/cod_produto nao encontrada na aba de layout")
    if idx_loc < 0 and (idx_galpao < 0 or idx_rua < 0 or idx_est < 0):
        raise RuntimeError("Colunas de endereco nao encontradas na aba de layout")

    out: Dict[str, Set[str]] = {}
    for _, row in df_layout.iterrows():
        cod = clean_text(row.iloc[idx_cod])
        if not cod:
            continue

        addrs: List[str] = []
        if idx_loc >= 0:
            addrs.extend(parse_endereco_cell(row.iloc[idx_loc]))

        if not addrs and idx_galpao >= 0 and idx_rua >= 0 and idx_est >= 0:
            esc = clean_text(row.iloc[idx_esc]) if idx_esc >= 0 else ""
            if not esc and idx_esc_num >= 0 and idx_nivel >= 0:
                esc_num = clean_text(row.iloc[idx_esc_num])
                nivel = clean_text(row.iloc[idx_nivel]).upper()
                if esc_num and nivel:
                    esc = f"{esc_num}{nivel}"
            addr = build_address_from_parts(
                row.iloc[idx_galpao], row.iloc[idx_rua], row.iloc[idx_est], esc
            )
            if addr:
                addrs.append(addr)

        for addr in addrs:
            n = normalize_addr(addr)
            if not is_valid_address(n):
                continue
            out.setdefault(cod, set()).add(n)

    return out


def build_address_from_parts(g, r, p, e):
    g = "" if g is None else str(g).strip()
    r = "" if r is None else str(r).strip()
    p = "" if p is None else str(p).strip()
    e = "" if e is None else str(e).strip()
    if not g or not r or not p or not e:
        return ""
    p = normalize_estante(p)
    return normalize_addr(f"{g}-{r}-{p}-{e}")


def norm_tipo(value: str) -> str:
    s = "" if value is None else str(value).strip().lower()
    s = s.replace("_", " ")
    s = re.sub(r"\s+", "", s)
    return s


def build_baixa_enderecos_antigos(
    new_addresses_by_cod: Dict[str, Set[str]],
    df_stock: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, int]]:
    if not new_addresses_by_cod:
        empty = pd.DataFrame(columns=BAIXA_COLUMNS)
        return empty, pd.DataFrame(columns=PROVA_COLUMNS), {
            "produtos_layout": 0,
            "linhas_estoque": 0,
            "linhas_baixa": 0,
            "falhas_prova": 0,
        }

    cols = list(df_stock.columns)
    idx_loc = find_col(cols, "location_id", "endereco", "endereco_generated")
    idx_cod = find_col(cols, "cod_produto", "product_code", "codigo")
    idx_galpao = find_col(cols, "galpao", "galpao_id")
    idx_rua = find_col(cols, "rua", "rua_num")
    idx_est = find_col(cols, "posicao_pallete", "posicao", "estante")
    idx_esc = find_col(cols, "escaninho_nivel", "escaninho")
    idx_qtd = find_col(cols, "quantidade", "quantity", "qtd")
    idx_validade = find_col(cols, "data_validade", "validade", "expiration_date")

    if idx_cod < 0 or idx_qtd < 0:
        raise RuntimeError("Colunas cod_produto/quantidade nao encontradas na aba Estoque_Atual")
    if idx_loc < 0 and (idx_galpao < 0 or idx_rua < 0 or idx_est < 0 or idx_esc < 0):
        raise RuntimeError("Colunas de endereco nao encontradas na aba Estoque_Atual")

    agg: Dict[Tuple[str, str, str, str, str, Any], float] = {}
    linhas_estoque_consideradas = 0

    for _, row in df_stock.iterrows():
        cod = clean_text(row.iloc[idx_cod])
        if not cod or cod not in new_addresses_by_cod:
            continue

        qtd = parse_number(row.iloc[idx_qtd])
        if qtd is None or qtd == 0:
            continue

        addr = row.iloc[idx_loc] if idx_loc >= 0 else ""
        if (not addr) and idx_galpao >= 0 and idx_rua >= 0 and idx_est >= 0 and idx_esc >= 0:
            addr = build_address_from_parts(row.iloc[idx_galpao], row.iloc[idx_rua], row.iloc[idx_est], row.iloc[idx_esc])
        parts = split_address_parts(addr)
        if not parts:
            continue

        validade = normalize_validade(row.iloc[idx_validade]) if idx_validade >= 0 else ""
        key = (cod, parts[0], parts[1], parts[2], parts[3], validade)
        agg[key] = agg.get(key, 0.0) + float(qtd)
        linhas_estoque_consideradas += 1

    out_rows = []
    prova_rows = []
    falhas_prova = 0
    for (cod, galpao, rua, estante, escaninho, validade), qtd_sum in agg.items():
        addr = normalize_addr(f"{galpao}-{rua}-{estante}-{escaninho}")
        novos = new_addresses_by_cod.get(cod, set())

        # Nunca dar baixa no endereco novo.
        if addr in novos:
            continue

        qtd_neg = -abs(qtd_sum)
        if qtd_neg == 0:
            continue

        estante_fmt = normalize_estante(estante)
        qtd_old = to_compact_number(qtd_sum)
        qtd_baixa = to_compact_number(qtd_neg)
        confere_qtd = qtd_baixa == to_compact_number(-abs(qtd_sum))
        confere_addr_antigo = addr not in novos
        if not confere_qtd or not confere_addr_antigo:
            falhas_prova += 1

        out_rows.append(
            [
                cod,
                galpao,
                rua,
                estante_fmt,
                escaninho,
                "DANO",
                qtd_baixa,
                "enderecamento",
                validade,
                "DESCARTE",
                "UN",
            ]
        )
        prova_rows.append(
            [
                cod,
                addr,
                " | ".join(sorted(novos)),
                qtd_old,
                qtd_baixa,
                validade,
                confere_qtd,
                confere_addr_antigo,
            ]
        )

    out_rows.sort(key=lambda r: (r[0], r[1], r[2], r[3], r[4], str(r[8])))
    out_df = pd.DataFrame(out_rows, columns=BAIXA_COLUMNS)
    prova_df = pd.DataFrame(prova_rows, columns=PROVA_COLUMNS)
    if not prova_df.empty:
        prova_df = prova_df.sort_values(
            by=["cod_produto", "endereco_antigo", "data_validade"],
            key=lambda c: c.astype(str),
            kind="stable",
        ).reset_index(drop=True)

    stats = {
        "produtos_layout": len(new_addresses_by_cod),
        "linhas_estoque": linhas_estoque_consideradas,
        "linhas_baixa": len(out_df),
        "falhas_prova": falhas_prova,
    }
    return out_df, prova_df, stats


def main():
    file_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_FILE

    # Carrega abas (lendo valores a partir das formulas do Google Sheets exportado)
    wb = load_workbook(file_path, data_only=False)
    df_locs = sheet_to_df(wb, SHEET_LOCS)
    df_end = sheet_to_df(wb, SHEET_END)
    df_block = sheet_to_df(wb, SHEET_BLOCK)
    df_cod = sheet_to_df(wb, SHEET_COD)
    df_layout, layout_sheet_used = load_layout_df(wb)
    if SHEET_STOCK in wb.sheetnames:
        df_stock = sheet_to_df(wb, SHEET_STOCK)
    else:
        df_stock = pd.DataFrame()
    new_addresses_by_cod = collect_new_addresses_by_product(df_layout)

    # Codigos de barras: cod_produto -> descricao
    cod_cols = list(df_cod.columns)
    idx_cod = find_col(cod_cols, "cod_produto", "product_code")
    idx_desc = find_col(cod_cols, "descricao", "descrição", "produto", "product_name", "nome")
    by_cod = {}
    if idx_cod >= 0 and idx_desc >= 0:
        for _, row in df_cod.iterrows():
            cod = "" if row.iloc[idx_cod] is None else str(row.iloc[idx_cod]).strip()
            desc = "" if row.iloc[idx_desc] is None else str(row.iloc[idx_desc]).strip()
            if cod and desc and cod not in by_cod:
                by_cod[cod] = desc

    # Enderecados (LocalCadastrado e LocalBipado)
    end_cols = list(df_end.columns)
    idx_cad = find_col(end_cols, "localcadastrado")
    idx_bip = find_col(end_cols, "localbipado")
    if idx_cad < 0 or idx_bip < 0:
        raise RuntimeError("Colunas LocalCadastrado/LocalBipado nao encontradas na aba Endereçamento")

    enderecados = set()
    for _, row in df_end.iterrows():
        cad = row.iloc[idx_cad] if idx_cad >= 0 else None
        bip = row.iloc[idx_bip] if idx_bip >= 0 else None
        for a in parse_endereco_cell(cad):
            enderecados.add(a)
        for a in parse_endereco_cell(bip):
            enderecados.add(a)

    # Bloqueados
    block_cols = list(df_block.columns)
    idx_addr = find_col(block_cols, "endereco_generated", "endereco", "location_id")
    idx_gond = find_col(block_cols, "is_gondola")
    idx_blocked = find_col(block_cols, "is_blocked")
    idx_galpao = find_col(block_cols, "galpao")
    idx_rua = find_col(block_cols, "rua")
    idx_pos = find_col(block_cols, "posicao_pallete", "posicao", "estante")
    idx_esc = find_col(block_cols, "escaninho_nivel", "escaninho")

    blocked = set()
    for _, row in df_block.iterrows():
        addr = ""
        if idx_addr >= 0:
            addr = row.iloc[idx_addr]
        if (not addr) and idx_galpao >= 0 and idx_rua >= 0 and idx_pos >= 0 and idx_esc >= 0:
            addr = build_address_from_parts(row.iloc[idx_galpao], row.iloc[idx_rua], row.iloc[idx_pos], row.iloc[idx_esc])
        addr = normalize_addr(addr)
        if not addr:
            continue
        should_block = False
        if idx_blocked >= 0:
            flag = parse_flag(row.iloc[idx_blocked])
            if flag == 1:
                should_block = True
        if not should_block and idx_gond >= 0:
            flag = parse_flag(row.iloc[idx_gond])
            if flag == 0:
                should_block = True
        if should_block:
            blocked.add(addr)

    # Localizações produtos
    loc_cols = list(df_locs.columns)
    idx_loc = find_col(loc_cols, "location_id", "endereco", "endereco_generated")
    idx_tipo = find_col(loc_cols, "tipo_equipamento_final", "tipo_equipamento")
    idx_pc = find_col(loc_cols, "product_code", "cod_produto")
    idx_nome = find_col(loc_cols, "product_name", "produto", "nome")
    if idx_loc < 0 or idx_pc < 0:
        raise RuntimeError("Colunas location_id/product_code nao encontradas na aba Localizações produtos")

    allowed = {"prateleira", "prateleiralateral", "prateleiraalta"}

    out_rows = []
    seen_addr = set()
    for _, row in df_locs.iterrows():
        addr = normalize_addr(row.iloc[idx_loc])
        if not addr:
            continue
        if addr in blocked:
            continue
        if addr in enderecados:
            continue

        if idx_tipo >= 0:
            tipo = norm_tipo(row.iloc[idx_tipo])
            if tipo and tipo not in allowed:
                continue

        cod = "" if row.iloc[idx_pc] is None else str(row.iloc[idx_pc]).strip()
        nome = "" if idx_nome < 0 or row.iloc[idx_nome] is None else str(row.iloc[idx_nome]).strip()
        if not nome and cod and cod in by_cod:
            nome = by_cod.get(cod, "")

        if addr in seen_addr:
            continue
        seen_addr.add(addr)
        out_rows.append([addr, nome, cod])

    out_df = pd.DataFrame(out_rows, columns=["endereco_bipar", "produto", "cod_produto"])
    if df_stock.empty:
        baixa_df = pd.DataFrame(columns=BAIXA_COLUMNS)
        prova_df = pd.DataFrame(columns=PROVA_COLUMNS)
        baixa_stats = {"produtos_layout": len(new_addresses_by_cod), "linhas_estoque": 0, "linhas_baixa": 0, "falhas_prova": 0}
        print(f"Aviso: aba {SHEET_STOCK} nao encontrada ou vazia. {OUT_SHEET_BAIXA} sera gerada vazia.")
    else:
        baixa_df, prova_df, baixa_stats = build_baixa_enderecos_antigos(new_addresses_by_cod, df_stock)

    # Escreve as abas de saida (substitui se existir)
    outputs = {
        OUT_SHEET: out_df,
        OUT_SHEET_BAIXA: baixa_df,
        OUT_SHEET_PROVA: prova_df,
    }
    try:
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            for sheet_name, df in outputs.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)
    except TypeError:
        wb = load_workbook(file_path)
        changed = False
        for sheet_name in outputs:
            if sheet_name in wb.sheetnames:
                wb.remove(wb[sheet_name])
                changed = True
        if changed:
            wb.save(file_path)
        with pd.ExcelWriter(file_path, engine="openpyxl", mode="a") as writer:
            for sheet_name, df in outputs.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

    # Validacao: nenhum endereco de saida pode estar enderecado
    out_set = set(out_df["endereco_bipar"].tolist())
    inter = out_set.intersection(enderecados)

    print("Arquivo:", file_path)
    print("Layout usado:", layout_sheet_used)
    print("Produtos no layout novo:", len(new_addresses_by_cod))
    print("Enderecados:", len(enderecados))
    print("Bloqueados:", len(blocked))
    print("Saida:", len(out_df))
    print("Baixa antigos:", len(baixa_df))
    print("Linhas estoque usadas:", baixa_stats.get("linhas_estoque", 0))
    print("Falhas prova baixa:", baixa_stats.get("falhas_prova", 0))
    print("Intersecao saida x enderecados:", len(inter))
    if inter:
        print("Exemplos em conflito:", list(sorted(inter))[:10])
        sys.exit(2)


if __name__ == "__main__":
    main()
