from copy import copy
from pathlib import Path
from collections import defaultdict

import openpyxl


INPUT_PATH = Path("/Users/andrelobo/Downloads/Desativar e liberar endereço - PAMPLONA (2).xlsx")
OUTPUT_PATH = Path("/Users/andrelobo/Downloads/Plano_Enderecamento_Final_JOIN_ATUALIZADO.xlsx")

SHEET_DEACTIVATE = "Desativar e liberar endereço -"
SHEET_READDRESS = "Novo Endereçamento R5 e R7 - G"
SHEET_JOIN = "Plano_Enderecamento_Final_JOIN"


LOCATION_COLUMNS = {
    "location_id",
    "galpao_id",
    "rua_num",
    "equipamento_num",
    "tipo_equipamento",
    "nivel",
    "escaninho_num_no_nivel",
    "capacidade_l",
    "tipo_equipamento_final",
    "is_hot_zone",
    "is_nivel_alto",
    "is_nivel_inferior",
    "slot_duplo",
}

PRODUCT_COLUMNS = {
    "product_code",
    "product_name",
    "quantidade",
    "curva",
    "grupo",
    "categoria_armazenagem",
    "vol_l_unitario",
    "vol_L_unitario",
    "venda_total",
    "nm_fabricante",
    "altura_cm",
    "peso_kg_unitario",
    "subcategoria",
    "is_pesado",
    "is_alto",
    "is_realocado",
    "produto_alocado_code",
    "grupo_alocado",
}


def normalize_code(value):
    if value is None:
        return None
    return str(value).strip().upper()


def split_destinations(value):
    return [part.strip() for part in str(value).split("+") if part and part.strip()]


def clone_cell(source_cell, target_cell):
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)


def copy_sheet_defaults(source_ws, target_ws):
    target_ws.sheet_format.defaultColWidth = source_ws.sheet_format.defaultColWidth
    target_ws.sheet_format.defaultRowHeight = source_ws.sheet_format.defaultRowHeight
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
    target_ws.page_margins = copy(source_ws.page_margins)
    target_ws.page_setup = copy(source_ws.page_setup)
    target_ws.print_options = copy(source_ws.print_options)
    target_ws.sheet_state = source_ws.sheet_state
    for key, dim in source_ws.column_dimensions.items():
        target_dim = target_ws.column_dimensions[key]
        target_dim.width = dim.width
        target_dim.hidden = dim.hidden
        target_dim.bestFit = dim.bestFit
        target_dim.outlineLevel = dim.outlineLevel
        target_dim.collapsed = dim.collapsed


def worksheet_index(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


def first_non_blank(records, column_name):
    for record in records:
        value = record["values"][column_name]
        if value not in (None, ""):
            return value
    return None


def build_new_row(product_records, template_record, instruction_name):
    values = dict(template_record["values"])
    for column_name in PRODUCT_COLUMNS:
        if column_name == "product_name":
            values[column_name] = first_non_blank(product_records, column_name) or instruction_name
        elif column_name == "is_realocado":
            current = first_non_blank(product_records, column_name)
            values[column_name] = True if current in (None, "") else current
        elif column_name in {"produto_alocado_code", "grupo_alocado"}:
            continue
        else:
            values[column_name] = first_non_blank(product_records, column_name)

    values["location_id_atual"] = values["location_id"]
    values["produto_alocado_code"] = values["product_code"]
    values["grupo_alocado"] = values["grupo"]
    if not values.get("product_name"):
        values["product_name"] = instruction_name
    return values


def main():
    source_wb = openpyxl.load_workbook(INPUT_PATH)
    ws_deactivate = source_wb[SHEET_DEACTIVATE]
    ws_readdress = source_wb[SHEET_READDRESS]
    ws_join = source_wb[SHEET_JOIN]

    deactivate_idx = worksheet_index(ws_deactivate)
    readdress_idx = worksheet_index(ws_readdress)
    join_idx = worksheet_index(ws_join)

    headers = [ws_join.cell(1, c).value for c in range(1, ws_join.max_column + 1)]

    deactivate_codes = {
        normalize_code(ws_deactivate.cell(r, deactivate_idx["product_code"]).value)
        for r in range(2, ws_deactivate.max_row + 1)
        if normalize_code(ws_deactivate.cell(r, deactivate_idx["product_code"]).value)
    }

    readdress_instructions = []
    for r in range(2, ws_readdress.max_row + 1):
        product_code = normalize_code(ws_readdress.cell(r, readdress_idx["product_code"]).value)
        new_dest = ws_readdress.cell(r, readdress_idx["ENDEREÇO NOVO"]).value
        if not product_code or not new_dest:
            continue
        readdress_instructions.append(
            {
                "row": r,
                "product_code": product_code,
                "product_name": ws_readdress.cell(r, readdress_idx["product_name"]).value,
                "destinations": split_destinations(new_dest),
            }
        )

    join_records = []
    join_by_product = defaultdict(list)
    join_by_location = defaultdict(list)
    header_row = ws_join[1]
    for r in range(2, ws_join.max_row + 1):
        values = {header: ws_join.cell(r, join_idx[header]).value for header in headers}
        record = {
            "source_row": r,
            "values": values,
            "product_code_norm": normalize_code(values["product_code"]),
            "location_id": values["location_id"],
        }
        join_records.append(record)
        join_by_product[record["product_code_norm"]].append(record)
        join_by_location[record["location_id"]].append(record)

    target_locations = {
        destination
        for instruction in readdress_instructions
        for destination in instruction["destinations"]
    }
    readdress_codes = {instruction["product_code"] for instruction in readdress_instructions}

    kept_records = [
        record
        for record in join_records
        if record["product_code_norm"] not in deactivate_codes
        and record["product_code_norm"] not in readdress_codes
        and record["location_id"] not in target_locations
    ]

    new_records = []
    template_use_count = defaultdict(int)
    for instruction in readdress_instructions:
        product_records = join_by_product[instruction["product_code"]]
        if not product_records:
            continue
        for destination in instruction["destinations"]:
            templates = join_by_location[destination]
            template_index = template_use_count[destination]
            if template_index >= len(templates):
                template_index = len(templates) - 1
            template_record = templates[template_index]
            template_use_count[destination] += 1

            new_values = build_new_row(product_records, template_record, instruction["product_name"])
            new_values["location_id"] = destination
            new_values["location_id_atual"] = destination

            new_records.append(
                {
                    "source_row": template_record["source_row"],
                    "style_row": template_record["source_row"],
                    "values": new_values,
                }
            )

    final_records = []
    for record in kept_records:
        final_records.append(
            {
                "source_row": record["source_row"],
                "style_row": record["source_row"],
                "values": record["values"],
            }
        )
    final_records.extend(new_records)

    final_records.sort(
        key=lambda item: (
            item["source_row"],
            normalize_code(item["values"]["product_code"]) or "",
            item["values"]["location_id"] or "",
        )
    )

    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    output_ws.title = SHEET_JOIN
    copy_sheet_defaults(ws_join, output_ws)

    for col in range(1, ws_join.max_column + 1):
        clone_cell(ws_join.cell(1, col), output_ws.cell(1, col))
    output_ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws_join.max_column)}1"

    out_row = 2
    for item in final_records:
        style_row = item["style_row"]
        src_dim = ws_join.row_dimensions[style_row]
        if src_dim.height is not None:
            output_ws.row_dimensions[out_row].height = src_dim.height
        for col_index, header in enumerate(headers, start=1):
            target_cell = output_ws.cell(out_row, col_index)
            clone_cell(ws_join.cell(style_row, col_index), target_cell)
            target_cell.value = item["values"].get(header)
        out_row += 1

    output_wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
