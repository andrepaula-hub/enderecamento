from copy import copy
from pathlib import Path

import openpyxl


INPUT_PATH = Path("/Users/andrelobo/Downloads/Pamplona (2).xlsx")
OUTPUT_PATH = Path("/Users/andrelobo/Downloads/Plano_Enderecamento_Final_JOIN.xlsx")
BACKUP_SHEET = "Plano_Enderecamento_Final_BACKU"
FINAL_SHEET = "Plano_Enderecamento_Final"
OUTPUT_SHEET = "Plano_Enderecamento_Final_JOIN"

SHELF_TYPES = {"prateleira", "prateleira_pamplona"}
COLD_TYPES = {"freezer", "geladeira"}


def clone_cell(source_cell, target_cell):
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.font:
        target_cell.font = copy(source_cell.font)
    if source_cell.fill:
        target_cell.fill = copy(source_cell.fill)
    if source_cell.border:
        target_cell.border = copy(source_cell.border)
    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)
    if source_cell.protection:
        target_cell.protection = copy(source_cell.protection)


def copy_sheet_defaults(source_ws, target_ws):
    target_ws.sheet_format.defaultColWidth = source_ws.sheet_format.defaultColWidth
    target_ws.sheet_format.defaultRowHeight = source_ws.sheet_format.defaultRowHeight
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
    target_ws.page_margins = copy(source_ws.page_margins)
    target_ws.page_setup = copy(source_ws.page_setup)
    target_ws.print_options = copy(source_ws.print_options)
    target_ws.sheet_state = source_ws.sheet_state
    for key, dim in source_ws.column_dimensions.items():
        target_dim = target_ws.column_dimensions[key]
        target_dim.width = dim.width
        target_dim.hidden = dim.hidden
        target_dim.bestFit = dim.bestFit
        target_dim.outlineLevel = dim.outlineLevel
        target_dim.collapsed = dim.collapsed


def build_groups(ws):
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {name: pos + 1 for pos, name in enumerate(header)}
    groups = {}
    order = []
    for row in range(2, ws.max_row + 1):
        location_id = ws.cell(row, idx["location_id"]).value
        if location_id not in groups:
            groups[location_id] = []
            order.append(location_id)
        groups[location_id].append(row)
    return header, idx, groups, order


def location_type_set(ws, groups, idx, location_id):
    return {
        ws.cell(row, idx["tipo_equipamento_final"]).value
        for row in groups.get(location_id, [])
    }


def main():
    source_wb = openpyxl.load_workbook(INPUT_PATH)
    backup_ws = source_wb[BACKUP_SHEET]
    final_ws = source_wb[FINAL_SHEET]

    header, backup_idx, backup_groups, backup_order = build_groups(backup_ws)
    _, final_idx, final_groups, _ = build_groups(final_ws)

    cold_locations_in_final = {
        location_id
        for location_id in final_groups
        if location_type_set(final_ws, final_groups, final_idx, location_id) & COLD_TYPES
    }

    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    output_ws.title = OUTPUT_SHEET
    copy_sheet_defaults(backup_ws, output_ws)

    for col in range(1, backup_ws.max_column + 1):
        clone_cell(backup_ws.cell(1, col), output_ws.cell(1, col))
    output_ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(backup_ws.max_column)}1"

    out_row = 2
    handled_locations = set()
    for location_id in backup_order:
        if location_id in handled_locations:
            continue

        if location_id in cold_locations_in_final:
            source_ws = final_ws
            row_numbers = final_groups[location_id]
        else:
            source_ws = backup_ws
            row_numbers = backup_groups[location_id]

        for row_number in row_numbers:
            src_dim = source_ws.row_dimensions[row_number]
            if src_dim.height is not None:
                output_ws.row_dimensions[out_row].height = src_dim.height
            for col in range(1, source_ws.max_column + 1):
                clone_cell(source_ws.cell(row_number, col), output_ws.cell(out_row, col))
            out_row += 1

        handled_locations.add(location_id)

    output_wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
