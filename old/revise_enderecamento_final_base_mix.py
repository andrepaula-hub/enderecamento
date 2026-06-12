from pathlib import Path
from collections import Counter

import openpyxl


BASE_WORKBOOK_PATH = Path("/Users/andrelobo/Downloads/enderecamentoFinal.xlsx")
RULES_WORKBOOK_PATH = Path("/Users/andrelobo/Downloads/Desativar e liberar endereço - PAMPLONA (2).xlsx")
FINAL_JOIN_PATH = Path("/Users/andrelobo/Downloads/Plano_Enderecamento_Final_JOIN_ATUALIZADO.xlsx")
OUTPUT_PATH = Path("/Users/andrelobo/Downloads/enderecamentoFinal_revisado.xlsx")

RULES_DEACTIVATE_SHEET = "Desativar e liberar endereço -"
RULES_READDRESS_SHEET = "Novo Endereçamento R5 e R7 - G"
RULES_JOIN_SHEET = "Plano_Enderecamento_Final_JOIN"
BASE_SHEET = "Base"
MIX_SHEET = "Mix"
FINAL_JOIN_SHEET = "Plano_Enderecamento_Final_JOIN"

TYPE_TO_COLUMN = {
    "freezer": "escaninhos_necessarios_freezer",
    "geladeira": "escaninhos_necessarios_geladeira",
    "geladeira_alta": "escaninhos_necessarios_geladeira_alta",
    "prateleira": "escaninhos_necessarios_prateleira",
    "prateleira_lateral": "escaninhos_necessarios_prateleira_lateral",
}

COUNT_COLUMNS = [
    "escaninhos_necessarios",
    "escaninhos_necessarios_freezer",
    "escaninhos_necessarios_geladeira",
    "escaninhos_necessarios_geladeira_alta",
    "escaninhos_necessarios_prateleira",
    "escaninhos_necessarios_prateleira_lateral",
]


def normalize_code(value):
    if value is None:
        return None
    return str(value).strip().upper()


def split_destinations(value):
    return [part.strip() for part in str(value).split("+") if part and part.strip()]


def worksheet_index(ws):
    return {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}


def infer_type_key(current_type, final_join_type):
    if current_type in TYPE_TO_COLUMN:
        return current_type
    if final_join_type == "freezer":
        return "freezer"
    if final_join_type == "geladeira":
        return "geladeira"
    if final_join_type in {"prateleira", "prateleira_pamplona"}:
        return "prateleira"
    return None


def main():
    wb_base = openpyxl.load_workbook(BASE_WORKBOOK_PATH)
    wb_rules = openpyxl.load_workbook(RULES_WORKBOOK_PATH, data_only=False)
    wb_final_join = openpyxl.load_workbook(FINAL_JOIN_PATH, data_only=False)

    ws_base = wb_base[BASE_SHEET]
    ws_mix = wb_base[MIX_SHEET]
    ws_deactivate = wb_rules[RULES_DEACTIVATE_SHEET]
    ws_readdress = wb_rules[RULES_READDRESS_SHEET]
    ws_original_join = wb_rules[RULES_JOIN_SHEET]
    ws_final_join = wb_final_join[FINAL_JOIN_SHEET]

    base_idx = worksheet_index(ws_base)
    mix_idx = worksheet_index(ws_mix)
    deactivate_idx = worksheet_index(ws_deactivate)
    readdress_idx = worksheet_index(ws_readdress)
    original_join_idx = worksheet_index(ws_original_join)
    final_join_idx = worksheet_index(ws_final_join)

    deactivate_codes = {
        normalize_code(ws_deactivate.cell(r, deactivate_idx["product_code"]).value)
        for r in range(2, ws_deactivate.max_row + 1)
        if normalize_code(ws_deactivate.cell(r, deactivate_idx["product_code"]).value)
    }

    readdress_codes = set()
    target_locations = set()
    for r in range(2, ws_readdress.max_row + 1):
        code = normalize_code(ws_readdress.cell(r, readdress_idx["product_code"]).value)
        if code:
            readdress_codes.add(code)
        target_locations.update(split_destinations(ws_readdress.cell(r, readdress_idx["ENDEREÇO NOVO"]).value))

    old_occupant_codes = set()
    for r in range(2, ws_original_join.max_row + 1):
        location_id = ws_original_join.cell(r, original_join_idx["location_id"]).value
        code = normalize_code(ws_original_join.cell(r, original_join_idx["product_code"]).value)
        if location_id in target_locations and code:
            old_occupant_codes.add(code)

    affected_codes = deactivate_codes | readdress_codes | old_occupant_codes

    final_slot_count = Counter()
    final_join_type = {}
    for r in range(2, ws_final_join.max_row + 1):
        code = normalize_code(ws_final_join.cell(r, final_join_idx["product_code"]).value)
        if not code:
            continue
        final_slot_count[code] += 1
        if code not in final_join_type:
            final_join_type[code] = ws_final_join.cell(r, final_join_idx["tipo_equipamento_final"]).value

    base_rows_to_delete = []
    for r in range(2, ws_base.max_row + 1):
        code = normalize_code(ws_base.cell(r, base_idx["product_code"]).value)
        if code not in affected_codes:
            continue

        new_total = final_slot_count.get(code, 0)
        current_type = ws_base.cell(r, base_idx["tipo_equipamento_base"]).value
        type_key = infer_type_key(current_type, final_join_type.get(code))

        for column_name in COUNT_COLUMNS:
            ws_base.cell(r, base_idx[column_name]).value = 0

        ws_base.cell(r, base_idx["escaninhos_necessarios"]).value = new_total
        if new_total > 0 and type_key:
            ws_base.cell(r, base_idx[TYPE_TO_COLUMN[type_key]]).value = new_total
            ws_base.cell(r, base_idx["tipo_equipamento_base"]).value = type_key
        elif new_total == 0:
            ws_base.cell(r, base_idx["tipo_equipamento_base"]).value = None

    for r in range(2, ws_mix.max_row + 1):
        code = normalize_code(ws_mix.cell(r, mix_idx["Código"]).value)
        if code in deactivate_codes:
            base_rows_to_delete.append(r)

    for row_number in reversed(base_rows_to_delete):
        ws_mix.delete_rows(row_number, 1)

    wb_base.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
